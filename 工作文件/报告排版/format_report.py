from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
src = root / '实验报告/三+四周实验报告.docx'
out = src.with_name('三+四周实验报告_排版版.docx')
d = Document(src)

def font(run, size, latin='Times New Roman', chinese='宋体', bold=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    f = run._r.get_or_add_rPr().get_or_add_rFonts()
    for a in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        f.attrib.pop(qn('w:'+a), None)
    f.set(qn('w:eastAsia'), chinese)
    if bold is not None:
        run.bold = bold

for s in d.sections:
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.3)

for i,p in enumerate(d.paragraphs):
    pf=p.paragraph_format
    pf.space_before=Pt(0); pf.space_after=Pt(6)
    pf.line_spacing=1.25
    pf.left_indent=pf.right_indent=pf.first_line_indent=Pt(0)
    pf.widow_control=True
    for r in p.runs: font(r, 12)
    if i in (0,11):
        p.style=d.styles['Heading 1']
        pf.space_before=Pt(0); pf.space_after=Pt(12)
        pf.keep_with_next=True
        pf.page_break_before=(i==11)
        for r in p.runs: font(r,18,chinese='黑体',bold=True)
    elif i in (1,6):
        p.style=d.styles['Heading 2']
        pf.space_after=Pt(8); pf.keep_with_next=True
        pf.page_break_before=(i==6)
        for r in p.runs: font(r,14,chinese='黑体',bold=True)
    elif p.text:
        pf.keep_with_next=True
        pf.space_before=Pt(6)
        for r in p.runs: font(r,12,bold=True)
    elif p._p.xpath('.//w:drawing'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing=1.0; pf.space_after=Pt(8)
    else:
        pf.space_after=Pt(0); pf.line_spacing=Pt(3)
        for r in p.runs: font(r,3)

for t in d.tables:
    t.autofit=False
    t.columns[0].width=Cm(16.4)
    pr=t._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is None: borders=OxmlElement('w:tblBorders'); pr.append(borders)
    for side in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    for row in t.rows:
        row.height=None
        trpr=row._tr.get_or_add_trPr()
        for e in list(trpr):
            if e.tag==qn('w:cantSplit'): trpr.remove(e)
        for c in row.cells:
            c.width=Cm(16.4)
            cp=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,v in [('top','80'),('bottom','80'),('left','110'),('right','110')]:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),v); el.set(qn('w:type'),'dxa'); margins.append(el)
            cp.append(margins)
            for j,p in enumerate(c.paragraphs):
                pf=p.paragraph_format
                pf.space_before=pf.space_after=Pt(0)
                pf.line_spacing=Pt(11)
                if not p.text.strip(): pf.line_spacing=Pt(5.5)
                pf.left_indent=pf.right_indent=pf.first_line_indent=Pt(0)
                pf.keep_with_next=False; pf.keep_together=True
                pf.widow_control=False
                pf.tab_stops.clear_all()
                for k in range(1,17): pf.tab_stops.add_tab_stop(Pt(k*24))
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs: font(r,10,latin='Consolas',bold=False)
                if j==0:
                    pf.space_before=Pt(5); pf.space_after=Pt(5); pf.keep_with_next=True
                    for r in p.runs: font(r,11,latin='Consolas',bold=True)
                if len(c.paragraphs)<45 and j<len(c.paragraphs)-1:
                    pf.keep_with_next=True

# Let long source listings paginate as paragraphs instead of oversized table rows.
for t in list(d.tables):
    for row in t.rows:
        for c in row.cells:
            for p in list(c.paragraphs):
                t._tbl.addprevious(p._p)
    t._tbl.getparent().remove(t._tbl)
for p in d.paragraphs:
    snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0')
    p._p.get_or_add_pPr().append(snap)
d.save(out)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(src) as a, ZipFile(out) as b:
    for part in a.namelist():
        if part.startswith('word/media/'):
            assert a.read(part)==b.read(part),part
        if part.startswith('word/') and part.endswith('.xml') and part in b.namelist():
            x,y=etree.fromstring(a.read(part)),etree.fromstring(b.read(part))
            assert x.xpath('//w:t/text()',namespaces=ns)==y.xpath('//w:t/text()',namespaces=ns),part
print('Saved; text and original image bytes unchanged:',out)
